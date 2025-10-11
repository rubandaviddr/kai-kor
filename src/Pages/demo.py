from pptx import Presentation
from pptx.util import Inches

# Create a presentation
prs = Presentation()

# Slide content
slides_content = [
    ("Introduction to React.js", "Building User Interfaces the Modern Way\nYour Name | Date | Organization"),
    ("What is React?", "• A JavaScript library for building user interfaces\n• Created by Facebook\n• Component-based architecture\n• Declarative and efficient"),
    ("Why Use React?", "• Reusable components\n• Virtual DOM for performance\n• Strong community and ecosystem\n• Used in Facebook, Instagram, Netflix, etc."),
    ("JSX – JavaScript + XML", "• Syntax extension for JavaScript\n• Looks like HTML, but it's JavaScript\n• Example:\n\nconst element = <h1>Hello, world!</h1>;"),
    ("Components in React", "• Functional vs Class Components\n• Building blocks of React apps\n• Props – data passed to components\n• Example of a functional component"),
    ("State & Events", "• useState for managing component state\n• Handling user events like clicks or input\n• Example: Counter App"),
    ("Lifecycle & useEffect", "• useEffect = run side effects (e.g., API calls)\n• Replaces lifecycle methods like componentDidMount\n• Example with data fetching"),
    ("React Router", "• For single-page app navigation\n• react-router-dom package\n• Routes, Route, Link components"),
    ("Building & Deployment", "• Use create-react-app\n• Dev server, hot reloading\n• Deployment options: Vercel, Netlify, GitHub Pages"),
    ("Q&A + Thank You", "• Open for questions\n• Share resources: React docs, tutorials, GitHub\n• Thank the audience"),
]

# Add slides
for title, content in slides_content:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = content

# Save presentation
prs.save("React_Presentation.pptx")
print("✅ React_Presentation.pptx has been created.")
